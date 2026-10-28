"""
AAG LIVE ALGO - Paper Trade (cloud version)

Logic:
  - Entry time par F&O ka top gainer -> uska ATM CALL kharido (paper)
  - Entry time par F&O ka top loser  -> uska ATM PUT kharido (paper)
  - Har leg ka alag SL / Target (rupees mein), ya time exit
  - Koi real order nahi jaata. Sirf Upstox ka live price padhte hain.

Secrets (GitHub Secrets se aate hain, code mein nahi likhne):
  UPSTOX_ACCESS_TOKEN
  TELEGRAM_TOKEN
  TELEGRAM_CHAT_ID   -> "apni_personal_id,channel_id"  (comma se alag; pehli ID = aapki)
Settings: config.json (phone se edit kar sakte ho)
"""
import csv
import gzip
import io
import json
import os
import sys
import time
from datetime import date, datetime, timedelta, timezone

import openpyxl
import pandas as pd
import requests
from openpyxl.styles import Alignment, Font, PatternFill

IST = timezone(timedelta(hours=5, minutes=30))
HERE = os.path.dirname(os.path.abspath(__file__))

UPSTOX_ACCESS_TOKEN = os.environ.get("UPSTOX_ACCESS_TOKEN", "")
TG_TOKEN = os.environ.get("TELEGRAM_TOKEN", "")
TG_CHAT = os.environ.get("TELEGRAM_CHAT_ID", "")
TEST_MODE = os.environ.get("TEST_MODE", "0") == "1"

UPSTOX_HEADERS = {
    "Accept": "application/json",
    "Authorization": f"Bearer {UPSTOX_ACCESS_TOKEN}",
}

DEFAULTS = {
    "enabled": True,
    "entry_time": "09:20",
    "exit_time": "12:30",
    "sl_amount": 2500,
    "target_amount": 4000,
    "num_lots": 1,
    "heartbeat_minutes": 10,
    "late_start_grace_minutes": 3,
    "list_size": 10,
}

BASE_URL = "https://www.nseindia.com"
GAINERS_URL = f"{BASE_URL}/api/live-analysis-variations?index=gainers"
LOSERS_URL = f"{BASE_URL}/api/live-analysis-variations?index=loosers"

_token_alert_sent = False


def now_ist():
    return datetime.now(IST)


def load_config():
    cfg = dict(DEFAULTS)
    try:
        with open(os.path.join(HERE, "config.json"), encoding="utf-8") as f:
            cfg.update(json.load(f))
    except Exception as e:
        print(f"config.json padh nahi paaye, default use ho rahe hain: {e}")
    return cfg


# =========================
# TELEGRAM
# Pehli chat ID = owner. Baaki (channel / friends) ko sirf trade messages jaate hain.
# Technical errors sirf owner ko jaate hain. Test mode mein bhi sirf owner ko.
# =========================

def chat_ids():
    return [c.strip() for c in TG_CHAT.split(",") if c.strip()]


def _send(ids, text):
    if not TG_TOKEN:
        return
    for c in ids:
        try:
            requests.post(
                f"https://api.telegram.org/bot{TG_TOKEN}/sendMessage",
                data={"chat_id": c, "text": text},
                timeout=20,
            )
        except Exception as e:
            print(f"Telegram error ({c}): {e}")


def tg(text):
    """Sabko bhejo (test mode mein sirf owner ko)."""
    print(text)
    ids = chat_ids()
    _send(ids[:1] if TEST_MODE else ids, text)


def tg_admin(text):
    """Sirf owner ko (technical baatein)."""
    print("[ADMIN]", text)
    _send(chat_ids()[:1], text)


def tg_file(path, caption=""):
    if not TG_TOKEN or not os.path.exists(path):
        return
    ids = chat_ids()
    for c in (ids[:1] if TEST_MODE else ids):
        try:
            with open(path, "rb") as f:
                requests.post(
                    f"https://api.telegram.org/bot{TG_TOKEN}/sendDocument",
                    data={"chat_id": c, "caption": caption},
                    files={"document": f},
                    timeout=60,
                )
        except Exception as e:
            print(f"Telegram file error ({c}): {e}")


def skip(reason):
    """Aaj trade nahi hua: owner ko wajah, baaki sabko simple line."""
    tg_admin(f"⚠️ Aaj trade skip: {reason}")
    if not TEST_MODE:
        tg("ℹ️ Aaj ka paper trade nahi hua.")


# =========================
# UPSTOX
# =========================

def upstox_get(url, params=None, tries=3):
    global _token_alert_sent
    for i in range(tries):
        try:
            r = requests.get(url, headers=UPSTOX_HEADERS, params=params, timeout=20)
            if r.status_code == 200:
                return r.json()
            print(f"Upstox {r.status_code}: {r.text[:200]}")
            if r.status_code in (401, 403):
                if not _token_alert_sent:
                    _token_alert_sent = True
                    tg_admin("❌ Upstox token mein problem (401/403). Naya token banakar GitHub Secret update karo.")
                return None
        except Exception as e:
            print(f"Upstox request error: {e}")
        time.sleep(2)
    return None


def get_live_ltp(instrument_key):
    if not instrument_key:
        return None
    data = upstox_get(
        "https://api.upstox.com/v2/market-quote/ltp",
        params={"instrument_key": instrument_key},
        tries=2,
    )
    try:
        vals = (data or {}).get("data", {})
        if vals:
            return float(list(vals.values())[0]["last_price"])
    except Exception as e:
        print(f"LTP parse error: {e}")
    return None


def is_market_holiday(d):
    data = upstox_get(f"https://api.upstox.com/v2/market/holidays/{d.isoformat()}", tries=1)
    try:
        for h in (data or {}).get("data", []):
            if h.get("holiday_type") == "TRADING_HOLIDAY" and "NSE" in (h.get("closed_exchanges") or []):
                return True, h.get("description", "")
    except Exception:
        pass
    return False, ""


def load_fo_instruments():
    """Lot sizes + stock instrument keys + F&O stock list."""
    print("\nF&O Instruments CSV loading...")
    try:
        r = requests.get(
            "https://assets.upstox.com/market-quote/instruments/exchange/complete.csv.gz",
            timeout=60,
        )
        df = pd.read_csv(io.BytesIO(gzip.decompress(r.content)))
        fo = df[(df["exchange"] == "NSE_FO") & (df["instrument_type"] == "OPTSTK")].copy()
        fo["base_symbol"] = fo["tradingsymbol"].str.extract(r"^([A-Z&\-]+)")
        lot_map = fo.groupby("base_symbol")["lot_size"].first().astype(int).to_dict()
        fo_syms = sorted(lot_map.keys())
        eq = df[(df["exchange"] == "NSE_EQ") & (df["instrument_type"] == "EQUITY")]
        eq_map = dict(zip(eq["tradingsymbol"], eq["instrument_key"]))
        print(f"  {len(lot_map)} F&O symbols loaded.")
        return lot_map, eq_map, fo_syms
    except Exception as e:
        print(f"  CSV load failed: {e}")
        return {}, {}, []


def get_eq_key(symbol, eq_map):
    if symbol in eq_map:
        return eq_map[symbol]
    data = upstox_get(
        "https://api.upstox.com/v2/instruments/search",
        params={"query": symbol, "exchange": "NSE_EQ"},
    )
    items = (data or {}).get("data", [])
    for inst in items:
        if (inst.get("trading_symbol", "").upper() == symbol.upper()
                and inst.get("exchange", "") == "NSE"):
            return inst["instrument_key"]
    for inst in items:
        if inst.get("instrument_key", "").startswith("NSE_EQ"):
            return inst["instrument_key"]
    return None


# =========================
# EXPIRY (monthly = last Tuesday)
# =========================

def last_tuesday(yr, mo):
    if mo == 12:
        last_day = date(yr + 1, 1, 1) - timedelta(days=1)
    else:
        last_day = date(yr, mo + 1, 1) - timedelta(days=1)
    while last_day.weekday() != 1:
        last_day -= timedelta(days=1)
    return last_day


def get_monthly_expiry(today):
    exp = last_tuesday(today.year, today.month)
    if exp < today:
        y, m = (today.year + 1, 1) if today.month == 12 else (today.year, today.month + 1)
        exp = last_tuesday(y, m)
    return exp


# =========================
# GAINERS / LOSERS
# =========================

def row_ltp(s):
    try:
        return float(s.get("lastPrice") or s.get("ltp") or s.get("closePrice") or 0)
    except Exception:
        return 0.0


def row_pct(s):
    try:
        return float(s.get("pChange") or s.get("perChange") or 0)
    except Exception:
        return 0.0


def fmt_list(title, rows, n):
    lines = [title]
    for i, s in enumerate(rows[:n], 1):
        lines.append(f"{i}. {s.get('symbol', '-')}  ₹{row_ltp(s):.2f}  ({row_pct(s):+.2f}%)")
    return "\n".join(lines)


def fetch_primary_gainers_losers():
    try:
        from curl_cffi import requests as cffi_requests

        s = cffi_requests.Session(impersonate="chrome124")
        s.get(BASE_URL, timeout=20)
        time.sleep(1)
        headers = {"Accept": "application/json", "Referer": BASE_URL}
        g = s.get(GAINERS_URL, headers=headers, timeout=20).json()
        l = s.get(LOSERS_URL, headers=headers, timeout=20).json()
        gainers = g.get("FOSec", {}).get("data", [])
        losers = l.get("FOSec", {}).get("data", [])
        if gainers and losers:
            return gainers, losers
    except Exception as e:
        print(f"primary fetch fail: {e}")
    return None


def upstox_gainers_losers(eq_map, fo_syms):
    """Pehla source fail ho toh Upstox ke quotes se top gainers/losers nikalo."""
    keys = [eq_map[s] for s in fo_syms if s in eq_map]
    rows = []
    for i in range(0, len(keys), 400):
        data = upstox_get(
            "https://api.upstox.com/v2/market-quote/quotes",
            params={"instrument_key": ",".join(keys[i:i + 400])},
        )
        for k, v in ((data or {}).get("data", {}) or {}).items():
            try:
                ltp = float(v.get("last_price"))
                chg = float(v.get("net_change"))
                prev = ltp - chg
                if prev > 0:
                    rows.append({
                        "symbol": v.get("symbol") or k.split(":")[-1],
                        "lastPrice": ltp,
                        "pChange": round(chg / prev * 100, 2),
                    })
            except Exception:
                continue
    if not rows:
        return None
    return (sorted(rows, key=lambda x: -x["pChange"])[:20],
            sorted(rows, key=lambda x: x["pChange"])[:20])


def get_gainers_losers(eq_map, fo_syms):
    res = fetch_primary_gainers_losers()
    if not res:
        print("Primary data nahi mila, Upstox se nikal rahe hain...")
        res = upstox_gainers_losers(eq_map, fo_syms)
    if not res:
        return None
    gainers, losers = res
    return (sorted(gainers, key=lambda x: -row_pct(x)),
            sorted(losers, key=lambda x: row_pct(x)))


# =========================
# OPTION CHAIN
# =========================

def get_atm_and_premium(symbol, ltp_fallback, expiry_date, lot_map, eq_map):
    eq_key = get_eq_key(symbol, eq_map)
    if not eq_key:
        return None
    chain, used_expiry = [], None
    # Agar expiry holiday par shift ho gayi ho toh pichhle din try karo
    for back in (0, 1, 2):
        exp = (expiry_date - timedelta(days=back)).strftime("%Y-%m-%d")
        data = upstox_get(
            "https://api.upstox.com/v2/option/chain",
            params={"instrument_key": eq_key, "expiry_date": exp},
        )
        chain = (data or {}).get("data", [])
        if chain:
            used_expiry = exp
            break
    if not chain:
        return None

    spot = float(chain[0].get("underlying_spot_price", ltp_fallback))
    strikes = [item["strike_price"] for item in chain]
    atm = min(strikes, key=lambda s: abs(s - spot))

    out = {"symbol": symbol, "atm": atm, "spot": spot, "expiry": used_expiry, "eq_key": eq_key,
           "lot": lot_map.get(symbol.upper()),
           "ce_key": None, "pe_key": None, "ce_ltp": None, "pe_ltp": None}
    for item in chain:
        if item["strike_price"] == atm:
            call = item.get("call_options", {}) or {}
            put = item.get("put_options", {}) or {}
            out["ce_key"] = call.get("instrument_key")
            out["ce_ltp"] = (call.get("market_data", {}) or {}).get("ltp")
            out["pe_key"] = put.get("instrument_key")
            out["pe_ltp"] = (put.get("market_data", {}) or {}).get("ltp")
            break
    return out


# =========================
# EXCEL + CSV LOG
# =========================

XL_COLS = ["Stock", "Type", "Strike", "Lot Size", "Lots", "Qty", "Entry Time", "Exit Time",
           "Spot Entry", "Spot Exit", "Entry (₹)", "Exit (₹)", "PnL (₹)", "Reason"]


def save_to_excel(trade_summary, mtm_logs, filename):
    wb = openpyxl.Workbook()
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    green_fill = PatternFill("solid", fgColor="C6EFCE")
    red_fill = PatternFill("solid", fgColor="FFC7CE")

    ws1 = wb.active
    ws1.title = "Trade Summary"
    ws1.append(XL_COLS)
    for cell in ws1[1]:
        cell.fill, cell.font = header_fill, header_font
        cell.alignment = Alignment(horizontal="center")
    pnl_col = XL_COLS.index("PnL (₹)") + 1
    for t in trade_summary:
        ws1.append([t[c] for c in XL_COLS])
        ws1.cell(ws1.max_row, pnl_col).fill = green_fill if t["PnL (₹)"] >= 0 else red_fill
    total = round(sum(t["PnL (₹)"] for t in trade_summary), 2)
    row = [""] * len(XL_COLS)
    row[pnl_col - 2] = "TOTAL PNL →"
    row[pnl_col - 1] = total
    ws1.append(row)
    ws1.cell(ws1.max_row, pnl_col).font = Font(bold=True)
    ws1.cell(ws1.max_row, pnl_col).fill = green_fill if total >= 0 else red_fill
    for col in ws1.columns:
        ws1.column_dimensions[col[0].column_letter].width = 15

    for title, log in mtm_logs:
        ws = wb.create_sheet(title)
        ws.append(["Time", "Option LTP (₹)", "MTM (₹)"])
        for cell in ws[1]:
            cell.fill, cell.font = header_fill, header_font
            cell.alignment = Alignment(horizontal="center")
        for r in log:
            ws.append([r["time"], r["ltp"], r["mtm"]])
            ws.cell(ws.max_row, 3).fill = green_fill if r["mtm"] >= 0 else red_fill
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 18

    wb.save(filename)
    print(f"Excel saved: {filename}")


def append_history(trade_summary, day):
    path = os.path.join(HERE, "trade_history.csv")
    new = not os.path.exists(path)
    with open(path, "a", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        if new:
            w.writerow(["date"] + XL_COLS)
        for t in trade_summary:
            w.writerow([day] + [t[c] for c in XL_COLS])


# =========================
# TRADE ENGINE (paper)
# =========================

def make_leg(direction, symbol, key, eq_key, strike, entry, spot, lot, lots, expiry, entry_time):
    return {"direction": direction, "opt": "CE" if direction == "CALL" else "PE",
            "symbol": symbol, "key": key, "eq_key": eq_key, "strike": strike,
            "entry": float(entry), "spot_entry": float(spot), "spot": float(spot),
            "spot_exit": None, "lot": int(lot), "lots": int(lots), "qty": int(lot) * int(lots),
            "expiry": expiry, "entry_time": entry_time,
            "status": "OPEN", "ltp": float(entry),
            "exit_price": None, "exit_time": None, "reason": None, "log": []}


def icon(lg):
    return "🟢" if lg["opt"] == "CE" else "🔴"


def leg_name(lg):
    return f"{lg['symbol']} {lg['strike']:g} {lg['opt']}"


def leg_mtm(lg):
    price = lg["ltp"] if lg["status"] == "OPEN" else lg["exit_price"]
    return round((price - lg["entry"]) * lg["qty"], 2)


def leg_entry_block(lg):
    return (f"{icon(lg)} BUY {lg['direction']}\n"
            f"Stock: {lg['symbol']}\n"
            f"Spot (CMP): ₹{lg['spot_entry']:.2f}\n"
            f"Strike: {lg['strike']:g}\n"
            f"Option type: {lg['direction']} ({lg['opt']})\n"
            f"Entry price: ₹{lg['entry']:.2f}\n"
            f"Lot size: {lg['lot']} × {lg['lots']} lot = Qty {lg['qty']}\n"
            f"Expiry: {lg['expiry']}")


def close_leg(lg, price, now_str, reason):
    lg["exit_price"], lg["exit_time"], lg["reason"], lg["status"] = price, now_str, reason, "CLOSED"
    spot = get_live_ltp(lg["eq_key"]) or lg["spot"]
    lg["spot_exit"] = spot
    tg(f"{reason}\n"
       f"{icon(lg)} {lg['direction']} — {leg_name(lg)}\n"
       f"Exit time: {now_str}\n"
       f"Option price: ₹{lg['entry']:.2f} → ₹{price:.2f}\n"
       f"Spot (CMP): ₹{lg['spot_entry']:.2f} → ₹{spot:.2f}\n"
       f"Lot size: {lg['lot']} × {lg['lots']} lot = Qty {lg['qty']}\n"
       f"PnL: ₹{leg_mtm(lg):+.2f}")


def live_pnl_msg(legs, now_str):
    lines = [f"💓 LIVE PNL — {now_str}"]
    total = 0.0
    for lg in legs:
        if lg["status"] == "OPEN":
            sp = get_live_ltp(lg["eq_key"])
            if sp:
                lg["spot"] = sp
            price, spot, st = lg["ltp"], lg["spot"], "OPEN"
        else:
            price, spot, st = lg["exit_price"], lg["spot_exit"] or lg["spot"], "CLOSED"
        m = leg_mtm(lg)
        total += m
        lines.append(f"\n{icon(lg)} {lg['direction']} {leg_name(lg)}  [{st}]\n"
                     f"Spot ₹{spot:.2f} | Option ₹{price:.2f}\n"
                     f"PnL ₹{m:+.2f}")
    lines.append(f"\n💰 Total PnL: ₹{total:+.2f}")
    return "\n".join(lines)


def run_trade(cfg, legs, exit_time):
    sl, target = cfg["sl_amount"], cfg["target_amount"]
    hb_secs = (1 if TEST_MODE else int(cfg["heartbeat_minutes"])) * 60
    last_hb = time.time()
    ce, pe = legs

    print(f"\n{'Time':<6} | {'CE LTP':>8} | {'CE MTM':>10} | {'PE LTP':>8} | {'PE MTM':>10} | {'Combined':>10}")
    print("-" * 70)

    while True:
        t0 = time.time()
        now_str = now_ist().strftime("%H:%M")
        past_exit = now_str >= exit_time or (not TEST_MODE and now_str > "15:30")

        if past_exit:
            for lg in legs:
                if lg["status"] == "OPEN":
                    ltp = get_live_ltp(lg["key"]) or lg["ltp"]
                    lg["ltp"] = ltp
                    close_leg(lg, ltp, now_str, "⏰ TIME EXIT")
            break
        if all(lg["status"] == "CLOSED" for lg in legs):
            break

        for lg in legs:
            if lg["status"] == "OPEN":
                ltp = get_live_ltp(lg["key"])
                if ltp:
                    lg["ltp"] = ltp
                lg["log"].append({"time": now_str, "ltp": lg["ltp"], "mtm": leg_mtm(lg)})

        combined = round(leg_mtm(ce) + leg_mtm(pe), 2)
        print(f"{now_str:<6} | ₹{ce['ltp']:>7.2f} | ₹{leg_mtm(ce):>9.2f} | "
              f"₹{pe['ltp']:>7.2f} | ₹{leg_mtm(pe):>9.2f} | ₹{combined:>9.2f}")

        for lg in legs:
            if lg["status"] == "OPEN":
                m = leg_mtm(lg)
                if m >= target:
                    close_leg(lg, lg["ltp"], now_str, "🎯 TARGET HIT")
                elif m <= -sl:
                    close_leg(lg, lg["ltp"], now_str, "🛑 SL HIT")

        if time.time() - last_hb >= hb_secs and any(lg["status"] == "OPEN" for lg in legs):
            last_hb = time.time()
            tg(live_pnl_msg(legs, now_str))

        time.sleep(max(1, 60 - (time.time() - t0)))


def wait_until_entry(entry_time, grace_min):
    h, m = map(int, entry_time.split(":"))
    entry_dt = now_ist().replace(hour=h, minute=m, second=0, microsecond=0)
    if now_ist() > entry_dt + timedelta(minutes=grace_min):
        return False
    last_print = 0
    while now_ist() < entry_dt:
        if time.time() - last_print > 300:
            print(f"Entry {entry_time} ka wait... abhi {now_ist().strftime('%H:%M:%S')} IST")
            last_print = time.time()
        time.sleep(5)
    return True


# =========================
# MAIN
# =========================

def main():
    cfg = load_config()
    today = now_ist().date()
    tag = "🧪 TEST " if TEST_MODE else ""

    if not UPSTOX_ACCESS_TOKEN:
        tg_admin("❌ UPSTOX_ACCESS_TOKEN secret nahi mila. GitHub Secrets check karo.")
        sys.exit(1)

    if not cfg["enabled"] and not TEST_MODE:
        tg_admin("⏸️ Algo config mein OFF hai (enabled=false). Aaj koi trade nahi.")
        return

    if not TEST_MODE:
        if today.weekday() >= 5:
            print("Weekend, exit.")
            return
        holiday, desc = is_market_holiday(today)
        if holiday:
            tg(f"🏖️ Aaj market holiday hai ({desc}). Paper trade nahi hoga.")
            return

    entry_time, exit_time = cfg["entry_time"], cfg["exit_time"]
    if TEST_MODE:
        n = now_ist()
        entry_time = n.strftime("%H:%M")
        exit_time = (n + timedelta(minutes=3)).strftime("%H:%M")

    tg(f"{tag}🟢 AAG Algo chalu (Paper trade)\n"
       f"Entry {entry_time} | Exit {exit_time}\n"
       f"Har leg: SL ₹{cfg['sl_amount']} | Target ₹{cfg['target_amount']} | Lots {cfg['num_lots']}")

    lot_map, eq_map, fo_syms = load_fo_instruments()
    if not lot_map:
        skip("F&O instruments list load nahi hui")
        return

    if not wait_until_entry(entry_time, int(cfg["late_start_grace_minutes"])):
        skip(f"algo late start hua (entry time {entry_time} nikal chuka), galat price se bachne ke liye skip")
        return

    res = get_gainers_losers(eq_map, fo_syms)
    if not res:
        skip("gainers/losers data nahi mila")
        return
    gainers, losers = res
    n = int(cfg["list_size"])
    tg(f"{tag}📊 {entry_time} — F&O TOP GAINERS & LOSERS\n\n"
       + fmt_list("🟢 TOP GAINERS", gainers, n) + "\n\n"
       + fmt_list("🔴 TOP LOSERS", losers, n))

    g, l = gainers[0], losers[0]
    g_stock, l_stock = str(g["symbol"]), str(l["symbol"])
    expiry = get_monthly_expiry(today)

    ce = get_atm_and_premium(g_stock, row_ltp(g), expiry, lot_map, eq_map)
    pe = get_atm_and_premium(l_stock, row_ltp(l), expiry, lot_map, eq_map)
    if not ce or not pe or not ce["ce_key"] or not pe["pe_key"] or not ce["lot"] or not pe["lot"]:
        skip(f"option chain/lot size nahi mila ({g_stock}/{l_stock}, expiry {expiry})")
        return

    ce_entry = get_live_ltp(ce["ce_key"]) or ce["ce_ltp"]
    pe_entry = get_live_ltp(pe["pe_key"]) or pe["pe_ltp"]
    if not ce_entry or not pe_entry:
        skip("entry price (LTP) nahi mila")
        return

    actual_entry = now_ist().strftime("%H:%M")
    lots = int(cfg["num_lots"])
    ce_leg = make_leg("CALL", g_stock, ce["ce_key"], ce["eq_key"], ce["atm"], ce_entry,
                      ce["spot"], ce["lot"], lots, ce["expiry"], actual_entry)
    pe_leg = make_leg("PUT", l_stock, pe["pe_key"], pe["eq_key"], pe["atm"], pe_entry,
                      pe["spot"], pe["lot"], lots, pe["expiry"], actual_entry)

    tg(f"{tag}🚀 TRADE ENTRY (Paper) — {actual_entry}\n\n"
       f"{leg_entry_block(ce_leg)}\n\n{leg_entry_block(pe_leg)}\n\n"
       f"SL ₹{cfg['sl_amount']} | Target ₹{cfg['target_amount']} (har leg)")

    legs = [ce_leg, pe_leg]
    run_trade(cfg, legs, exit_time)

    summary = []
    for lg in legs:
        summary.append({
            "Stock": lg["symbol"], "Type": lg["direction"], "Strike": f"{lg['strike']:g} {lg['opt']}",
            "Lot Size": lg["lot"], "Lots": lg["lots"], "Qty": lg["qty"],
            "Entry Time": lg["entry_time"], "Exit Time": lg["exit_time"],
            "Spot Entry": round(lg["spot_entry"], 2), "Spot Exit": round(lg["spot_exit"] or lg["spot"], 2),
            "Entry (₹)": lg["entry"], "Exit (₹)": lg["exit_price"], "PnL (₹)": leg_mtm(lg),
            "Reason": lg["reason"],
        })
    total = round(sum(t["PnL (₹)"] for t in summary), 2)

    lines = [f"{tag}📊 DAY SUMMARY (Paper) — {today}"]
    for t, lg in zip(summary, legs):
        lines.append(f"\n{icon(lg)} {t['Type']} {leg_name(lg)}\n"
                     f"Option: ₹{t['Entry (₹)']:.2f} → ₹{t['Exit (₹)']:.2f}\n"
                     f"Spot: ₹{t['Spot Entry']:.2f} → ₹{t['Spot Exit']:.2f}\n"
                     f"Lot {t['Lot Size']} × {t['Lots']} = Qty {t['Qty']}\n"
                     f"PnL ₹{t['PnL (₹)']:+.2f}  [{t['Reason']}]")
    lines.append(f"\n💰 TOTAL PnL: ₹{total:+.2f}")
    tg("\n".join(lines))

    xlsx = os.path.join(HERE, f"live_trade_log_{today}.xlsx")
    save_to_excel(summary, [("CE MTM Log", ce_leg["log"]), ("PE MTM Log", pe_leg["log"])], xlsx)
    tg_file(xlsx, f"Trade log {today}")
    if not TEST_MODE:
        append_history(summary, str(today))


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        tg_admin(f"❌ Algo crash hua: {type(e).__name__}: {e}")
        raise
